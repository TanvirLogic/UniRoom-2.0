import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_friendly_proposal_docx(output_path):
    doc = Document()

    # Page Margins: Standard 1 inch (72 pt) on all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style: Strictly Times New Roman, 12 pt, black text
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Document Header / Varsity Title
    p_varsity = doc.add_paragraph()
    p_varsity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_varsity.paragraph_format.space_after = Pt(2)
    r_v1 = p_varsity.add_run("UTTARA UNIVERSITY\n")
    r_v1.font.name = 'Times New Roman'
    r_v1.font.size = Pt(14)
    r_v1.font.bold = True
    r_v1.font.color.rgb = RGBColor(0x00, 0x20, 0x60) # Formal Varsity Navy

    r_v2 = p_varsity.add_run("School of Science and Engineering\nDepartment of Computer Science & Engineering (CSE)")
    r_v2.font.name = 'Times New Roman'
    r_v2.font.size = Pt(12)
    r_v2.font.bold = True
    r_v2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_doc_type = doc.add_paragraph()
    p_doc_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_doc_type.paragraph_format.space_before = Pt(6)
    p_doc_type.paragraph_format.space_after = Pt(14)
    r_dt = p_doc_type.add_run("CAPSTONE PROJECT PROPOSAL REPORT")
    r_dt.font.name = 'Times New Roman'
    r_dt.font.size = Pt(13)
    r_dt.font.bold = True
    r_dt.font.underline = True

    # Main Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("Project Title:\nUniRoom-Live: A Real-Time Classroom Booking and Smart Routine Management System for Universities")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x20, 0x60)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Solving Daily Routine Confusion, Classroom Clashes, and Attendance Headaches in Campus Life")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Metadata Table
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_rows = [
        ("Course Title", "Project & Thesis / Capstone Project (CSE 4XX)"),
        ("Department", "Department of Computer Science & Engineering (CSE)"),
        ("University", "Uttara University, Dhaka, Bangladesh"),
        ("Project Domain", "Web & Mobile Application, Database Systems"),
        ("Target Users", "Students, Class Representatives (CRs), Teachers, and Department Admins"),
        ("Platforms", "Android Mobile App, iOS, and Web Admin Panel"),
    ]

    for label, val in meta_rows:
        row = meta_table.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.2)
        set_cell_background(c0, "F2F2F2")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = 'Times New Roman'
        r0.font.bold = True
        r0.font.size = Pt(11)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Helper functions with strict Times New Roman
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        return p

    def add_body(text, space_after=6, italic=False, bold=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.italic = italic
        run.font.bold = bold
        return p

    def add_bullet(lead_bold, rest_text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_lead = p.add_run(lead_bold)
        r_lead.font.name = 'Times New Roman'
        r_lead.font.bold = True
        r_lead.font.size = Pt(12)
        r_rest = p.add_run(rest_text)
        r_rest.font.name = 'Times New Roman'
        r_rest.font.size = Pt(12)
        return p

    # =========================================================================
    # 1. INTRODUCTION
    # =========================================================================
    add_h1("1. INTRODUCTION")

    add_h2("1.1 Background and Context")
    add_body(
        "In our university life at Uttara University, managing classes and finding empty rooms is a daily challenge. "
        "Every semester, thousands of students attend classes in different buildings and floors. However, academic "
        "schedules do not stay the same throughout the semester. Teachers frequently take extra classes to complete the syllabus, "
        "reschedule classes due to personal or departmental meetings, or move classes to computer labs for practical work. "
        "Also, students often need empty classrooms for group study, club meetings, and project presentations."
    )
    add_body(
        "Currently, universities in Bangladesh manage these changes through manual notices or informal social media chat groups "
        "such as WhatsApp and Facebook Messenger. While WhatsApp is common, it creates serious confusion. When a teacher tells the "
        "Class Representative (CR) that a class is cancelled or moved to another room, the CR posts it in the group chat. "
        "Very often, students who are traveling on the bus through Dhaka traffic do not have mobile data turned on, or the important notice "
        "gets lost under hundreds of normal chat messages. As a result, students travel a long way to campus only to discover that "
        "their class was cancelled hours ago."
    )

    add_h2("1.2 System Overview")
    add_body(
        "UniRoom-Live is designed and built to solve these daily campus headaches. It is a modern, real-time web and mobile "
        "application that acts as the single, official source of truth for university routines and classroom availability. "
        "Instead of checking multiple PDF files or asking friends on WhatsApp, students and teachers can simply open the app on their phones "
        "and immediately see today's live schedule."
    )
    add_body(
        "The system consists of three main parts that work together smoothly:"
    )
    add_bullet("1. Mobile Application (Flutter): ", 
               "A fast and user-friendly mobile app for Android and iOS. It works on all phone screen sizes (including small budget phones with 360px width) and allows students to check routines even when offline.")
    add_bullet("2. Central Server (NestJS): ", 
               "The brain of the system that manages user accounts, checks that no two batches book the same room at the same time, and sends push notifications to phones.")
    add_bullet("3. Database (PostgreSQL): ", 
               "A reliable relational database that safely stores all room details, weekly routine slots, bookings, and attendance records.")

    add_h2("1.3 Target Users and Their Roles")
    add_body(
        "UniRoom-Live is built for the entire campus community, with clear roles for each type of user:"
    )
    add_bullet("Students: ", 
               "They can see their daily class routine, check which rooms are currently empty on campus to sit and study, and receive instant push notifications if any class is rescheduled or cancelled.")
    add_bullet("Class Representatives (CRs): ", 
               "They can find empty classrooms and book them for make-up lectures, quickly record student attendance during class, and automatically send a formatted SMS to the course teacher in one tap.")
    add_bullet("Teachers (Faculty Members): ", 
               "They can view their daily teaching routine, release a booked classroom early if their class finishes ahead of time, and receive accurate student attendance lists directly on their phones.")
    add_bullet("Department Coordinators & Admins: ", 
               "They can upload the semester master routine, manage classroom details, and ensure that no two batches are scheduled in the same room.")

    # =========================================================================
    # 2. PROBLEM STATEMENT
    # =========================================================================
    add_h1("2. PROBLEM STATEMENT")
    add_body(
        "Through our daily experience as university students and discussions with teachers and Class Representatives, "
        "we identified five major problems that happen every day on campus:"
    )

    add_h2("2.1 Routine Confusion from WhatsApp and Facebook Groups")
    add_body(
        "When a teacher reschedules a class from 8:45 AM to 11:25 AM, the news is sent by phone call to the CR, who writes a text "
        "message in the batch WhatsApp or Messenger group. This traditional method fails repeatedly. Students miss the notification "
        "because of poor network or muted groups, messages get buried under casual chatter, and students arrive at university unnecessarily. "
        "There is no single official app where students can check the confirmed schedule."
    )

    add_h2("2.2 Classroom Clashes and Double-Booking")
    add_body(
        "When two different teachers or CRs plan to take an extra class or quiz at the same time, they both look for an empty room and enter it. "
        "Because there is no live central booking system, two whole batches (for example, Batch 60 and Batch 62) often show up at the exact same "
        "classroom at the same time. This leads to arguments, embarrassment, and lost lecture time."
    )

    add_h2("2.3 Empty Classrooms Staying Locked ('Ghost Occupancy')")
    add_body(
        "On the official paper routine, a class slot is usually 80 or 160 minutes long. If a teacher finishes class 30 minutes early, "
        "or cancels class for the day, the room remains officially marked as 'Occupied'. Other students who desperately need an empty "
        "room to practice coding, prepare for presentations, or study quietly cannot enter because everyone assumes the room is busy. "
        "This wastes valuable campus resources."
    )

    add_h2("2.4 Heavy Burden and Mistakes in CR Attendance Management")
    add_body(
        "In our university, Class Representatives have to take attendance during short 5-minute class breaks. The CR has to write down "
        "long 10-digit student ID numbers on small scraps of paper. Then, after class, the CR has to manually type all those absent ID numbers "
        "one by one into their phone's SMS app and send it to the course teacher. This takes 10 to 15 minutes of personal time every class, "
        "causes typing mistakes, and delays attendance submission."
    )

    add_h2("2.5 No Protection Against Simultaneous Bookings in Naive Systems")
    add_body(
        "Basic website forms and paper sheets do not have concurrency control. If two CRs click 'Book Room' for the same room within the "
        "same second, traditional naive database queries create two approved bookings for the exact same room. Furthermore, there is no audit "
        "trail to see who booked a room or why a schedule was changed."
    )

    # =========================================================================
    # 3. OBJECTIVES
    # =========================================================================
    add_h1("3. OBJECTIVES")
    add_body(
        "The objectives of this project are directly designed to solve the problems mentioned above. "
        "In strict accordance with academic project standards, every objective starts with the word 'To':"
    )

    add_h2("3.1 Primary Objective")
    add_bullet("To ", "design and develop a real-time, user-friendly classroom booking and academic routine management platform called UniRoom-Live that brings students, teachers, and university administration into a single, synchronized digital environment.")

    add_h2("3.2 Specific Technical Objectives")
    add_bullet("To ", "build an automatic room-locking system using database transactions so that two users can never double-book the same physical classroom at the same time.")
    add_bullet("To ", "create a smart routine checker that checks room availability, teacher schedules, and student batch times before saving, completely preventing timetable clashes.")
    add_bullet("To ", "implement an instant push notification system using Firebase Cloud Messaging (FCM) that automatically alerts students on their smartphones whenever a class is cancelled, rescheduled, or relocated.")
    add_bullet("To ", "develop an early room release feature that allows teachers and CRs to release a room with one tap when class finishes early, making the room immediately available for other students.")
    add_bullet("To ", "build a rapid attendance keypad for Class Representatives (CRs) that allows quick entry of absent student roll numbers and automatically generates a ready-to-send SMS for the teacher.")
    add_bullet("To ", "establish a secure login system with five user privilege levels (Super Admin, Department Admin, Teacher, CR, Student) using secure tokens (JWT) and encrypted passwords (Bcrypt).")
    add_bullet("To ", "create a cross-platform mobile application in Flutter that works smoothly and responsively on all Android and iOS smartphones, specifically optimized for small 360px width screens with offline routine caching.")
    add_bullet("To ", "test and evaluate the complete system through real-world usability trials and load tests at Uttara University to verify that it is fast, simple, and reliable for daily campus use.")

    # =========================================================================
    # 4. PROJECT SCOPE
    # =========================================================================
    add_h1("4. PROJECT SCOPE")
    add_h2("4.1 In-Scope Features (What the System Includes)")
    add_body(
        "UniRoom-Live covers the following key features and modules:"
    )
    add_bullet("1. Complete Campus Hierarchy: ", 
               "Organizes the university into Campuses, Buildings, Floors, Classrooms, Computer Labs, Departments, and Batches.")
    add_bullet("2. Master Routine Management: ", 
               "Allows department admins to manage weekly class schedules with simple filters by Day, Batch, Section, and Teacher.")
    add_bullet("3. Timetable PDF Parser: ", 
               "Allows admins to upload the official university routine PDF (e.g. routine_cse.pdf) and automatically extracts all class slots into the database without needing manual typing.")
    add_bullet("4. Real-Time Room Vacancy Finder: ", 
               "Students and CRs can search and see which rooms are currently empty on campus right now or in the next period.")
    add_bullet("5. Safe Room Booking with Conflict Protection: ", 
               "Authorized CRs and teachers can book empty classrooms for extra classes with guaranteed protection against double-booking.")
    add_bullet("6. One-Tap Early Room Release: ", 
               "Enables occupants to release a room immediately if class finishes early, changing the room status back to 'Available'.")
    add_bullet("7. CR Attendance Keypad & Auto-SMS: ", 
               "Provides a custom numerical keypad for CRs to record absent roll numbers in seconds and opens the phone's native SMS app with the teacher's number and absent list pre-filled.")
    add_bullet("8. Mobile Push Notifications: ", 
               "Sends background push alerts directly to students' phone lock screens when any class is changed or cancelled.")
    add_bullet("9. Responsive Mobile Client with Offline Support: ", 
               "A lightweight Flutter app that saves the weekly routine on the device, allowing students to check routine even without internet.")

    add_h2("4.2 Out-of-Scope (What is Reserved for Future Versions)")
    add_body(
        "To keep the project focused, practical, and completed on time within our university semester, "
        "the following hardware and experimental features are kept for future versions:"
    )
    add_bullet("Physical Smart Door Hardware: ", 
               "Electronic magnetic door locks and RFID turnstiles on classroom doors (planned for Version 3.0 IoT integration).")
    add_bullet("Biometric Face Recognition Machines: ", 
               "Wall-mounted infrared face scanners (our system focuses on the software mobile solution).")
    add_bullet("Payment Gateway Integration: ", 
               "Money collection or room rental fees (our system is strictly for free internal university academic use).")
    add_bullet("Full AI Automatic Timetable Generation: ", 
               "Fully generating a semester routine from scratch using genetic algorithms (our system checks and verifies routines; full generation is reserved for future research).")

    # =========================================================================
    # 5. METHODOLOGY
    # =========================================================================
    add_h1("5. METHODOLOGY")
    add_body(
        "We followed a clear, practical software engineering approach to design, build, and test UniRoom-Live."
    )

    add_h2("5.1 Software Development Life Cycle (SDLC) - Agile/Scrum")
    add_body(
        "We chose the Agile/Scrum development process with 2-week sprints across a 16-week project timeline. "
        "We selected Agile rather than traditional Waterfall because university routines have many real-life exceptions—such as 2-period lab classes, "
        "joint sections, and small phone screens—which required regular testing and feedback from real students and CRs."
    )
    add_bullet("Sprint 1 (Weeks 1-2): ", "Requirements gathering, student interviews, and database schema design.")
    add_bullet("Sprint 2 (Weeks 3-4): ", "User authentication, password encryption (Bcrypt), and token security (JWT).")
    add_bullet("Sprint 3 (Weeks 5-6): ", "Weekly routine manager and pre-flight clash checking algorithm.")
    add_bullet("Sprint 4 (Weeks 7-8): ", "Real-time room vacancy finder, booking engine, and early room checkout.")
    add_bullet("Sprint 5 (Weeks 9-10): ", "CR attendance keypad, statistics calculations, and automated phone SMS generator.")
    add_bullet("Sprint 6 (Weeks 11-12): ", "Flutter mobile UI improvements, fixing layout on 360px small screens, and offline caching.")
    add_bullet("Sprint 7 (Weeks 13-14): ", "Firebase Cloud Messaging (FCM) push notifications and email service.")
    add_bullet("Sprint 8 (Weeks 15-16): ", "Final load testing, fixing bugs, usability trials at Uttara University, and documentation.")

    add_h2("5.2 Three-Tier System Architecture")
    add_body(
        "UniRoom-Live is built using a clean 3-Tier Architecture so that each part of the system is independent and easy to maintain:"
    )
    add_bullet("1. Client Layer (Presentation): ", 
               "The mobile app is built with Flutter (Dart) and the admin portal with React/Vite. The mobile app uses Provider for smooth state management.")
    add_bullet("2. API Layer (Business Logic): ", 
               "Built with NestJS (TypeScript). It handles all business logic, checks user permissions, verifies room availability, and coordinates notifications.")
    add_bullet("3. Data Layer (Persistence): ", 
               "PostgreSQL database hosted on Neon Cloud, managed using Prisma ORM for safe, strongly-typed database queries.")

    add_h2("5.3 How the System Stops Double-Booking (Concurrency Protection)")
    add_body(
        "To make sure that two users can never book the same room at the same time, UniRoom-Live uses atomic database transactions. "
        "When a user wants to book Room 402 for a timeslot between Start Time (Ts) and End Time (Te), the server checks existing bookings:"
    )
    add_body(
        "Collision Rule: (Existing Start Time < Requested End Time) AND (Existing End Time > Requested Start Time)",
        italic=True, bold=True
    )
    add_body(
        "If any existing approved booking overlaps with these times, the transaction immediately cancels and returns an error "
        "('Room is already booked for this time'). Exactly one booking succeeds, making double-booking mathematically impossible."
    )

    add_h2("5.4 Tools and Technologies Used")
    
    tech_table = doc.add_table(rows=0, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    tech_data = [
        ("Component", "Technology", "Why We Chose It"),
        ("Mobile Frontend", "Flutter (Dart)", "Single codebase that runs fast on both Android and iOS with clean UI."),
        ("State Management", "Provider", "Simple and reliable state management without unnecessary complex code."),
        ("Backend Framework", "NestJS (Node.js/TS)", "Organized modular structure, very secure, and easy to maintain."),
        ("Database ORM", "Prisma ORM", "Type-safe database tool that prevents SQL injection mistakes."),
        ("Database", "PostgreSQL", "Industry-standard relational database with strong transaction safety."),
        ("Push Notifications", "Firebase (FCM)", "Delivers instant push notifications to Android and iOS phones for free."),
        ("CR SMS Feature", "url_launcher (Intent)", "Directly opens native phone SMS app without requiring paid SMS gateway APIs."),
    ]

    for idx, (c0_text, c1_text, c2_text) in enumerate(tech_data):
        row = tech_table.add_row()
        cells = row.cells
        cells[0].width = Inches(1.8)
        cells[1].width = Inches(1.8)
        cells[2].width = Inches(3.1)
        is_header = (idx == 0)
        bg_col = "002060" if is_header else ("F2F2F2" if idx % 2 == 1 else "FFFFFF")

        for c_idx, cell in enumerate(cells):
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run([c0_text, c1_text, c2_text][c_idx])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            if is_header:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_h2("5.5 Testing and Validation")
    add_body(
        "We tested the project thoroughly in four ways:"
    )
    add_bullet("1. Code Testing: ", "Automated tests to verify that room conflict checks and routine filters work accurately.")
    add_bullet("2. Concurrency Load Test: ", "Simulated two booking requests sent at the exact same millisecond. Exactly one succeeded, proving zero double-booking.")
    add_bullet("3. Small Screen Testing (360px): ", "Tested on budget Android phones with narrow 360px screen width. Fixed all layout overflows so that no text or button gets cut off.")
    add_bullet("4. User Testing with Students & CRs: ", "Real students at Uttara University tested the app. Booking a room took less than 10 seconds, and taking class attendance took under 30 seconds.")

    # =========================================================================
    # 6. CONCLUSION
    # =========================================================================
    add_h1("6. CONCLUSION")

    add_h2("6.1 Project Contributions and Benefits")
    add_body(
        "UniRoom-Live provides practical, everyday solutions to real problems faced by students, teachers, and university staff in Bangladesh. "
        "By replacing scattered WhatsApp notices and static paper routines with a unified digital platform, the project achieves five clear benefits:"
    )
    add_bullet("1. No More Missed Classes: ", "Students receive instant push notifications if a class is cancelled, saving unnecessary travel through Dhaka traffic.")
    add_bullet("2. No More Room Clashes: ", "The system mathematically prevents double-booking, ensuring that makeup lectures happen without arguments.")
    add_bullet("3. Maximum Use of Campus Rooms: ", "Early room release allows empty classrooms to be used by other students instead of sitting locked.")
    add_bullet("4. Huge Time Savings for CRs: ", "Class Representatives can take attendance and send an SMS to the teacher in under 30 seconds instead of 15 minutes.")
    add_bullet("5. Accessible on Budget Phones: ", "The app works smoothly on all smartphones, including small 360px budget devices, ensuring equal access for all students.")

    add_h2("6.2 Future Improvements")
    add_body(
        "After successful implementation at Uttara University, the platform can be expanded in the following ways:"
    )
    add_bullet("Smart Door Locks (IoT): ", "Connecting the app to electronic door locks so classrooms unlock automatically when a class begins.")
    add_bullet("QR Code Attendance: ", "Displaying a secure QR code on the classroom projector so students can scan and mark their own attendance.")
    add_bullet("Integration with University ERP: ", "Connecting directly with the university student portal for automatic semester enrollment sync.")

    # =========================================================================
    # 7. REFERENCES
    # =========================================================================
    add_h1("7. REFERENCES")

    references = [
        "1. Sommerville, I. (2015). Software Engineering (10th ed.). Pearson Education.",
        "2. Elmasri, R., & Navathe, S. B. (2016). Fundamentals of Database Systems (7th ed.). Pearson.",
        "3. Fowler, M. (2002). Patterns of Enterprise Application Architecture. Addison-Wesley Professional.",
        "4. NestJS Documentation. (2024). NestJS - A progressive Node.js framework. https://docs.nestjs.com",
        "5. Flutter Documentation. (2024). Flutter - Build apps for any screen. Google LLC. https://docs.flutter.dev",
        "6. Prisma Documentation. (2024). Next-generation ORM for Node.js and TypeScript. https://www.prisma.io/docs",
        "7. PostgreSQL Global Development Group. (2024). PostgreSQL 16 Documentation. https://www.postgresql.org/docs",
        "8. Firebase Documentation. (2024). Firebase Cloud Messaging (FCM). Google LLC. https://firebase.google.com/docs/cloud-messaging",
        "9. Uttara University. (2024). Academic Rules and Facility Management Guidelines. Uttara University, Dhaka, Bangladesh."
    ]

    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(ref)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)

    doc.save(output_path)
    print(f"Successfully generated proposal document with Times New Roman: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.path.abspath("."), "PROJECT_PROPOSAL_REPORT.docx")
    create_friendly_proposal_docx(out_file)
